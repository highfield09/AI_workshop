"""Read-only previews for the OCR lesson; extraction remains the student's task."""

import base64
from html import escape
from itertools import islice
import os
from zipfile import BadZipFile

from IPython.display import HTML, display

from llm_workshop.course import ROOT
from llm_workshop.presentation import readable_html


FIRST_IMAGE = ROOT / "_for_STUDENT/data/notebook6/batch1_1/batch1-0001.jpg"
SPREADSHEET = ROOT / "_for_STUDENT/outputs/06_receipt_ocr/first_receipt.xlsx"


def show_receipt_preview(image_path=FIRST_IMAGE):
    if not image_path.is_file():
        return display(HTML(readable_html(
            "<b>Download the lesson data first.</b> In the terminal, run "
            "<code>python _for_TRAINER/scripts/prepare_receipt_data.py</code>, then rerun this cell."
        )))
    encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
    return display(HTML(readable_html(
        f"<b>{escape(image_path.name)}</b> · first image by filename, not first CSV row"
        f"<img src='data:image/jpeg;base64,{encoded}' alt='First invoice from the Kaggle dataset' "
        "style='display:block;width:100%;max-width:440px;height:auto;margin:12px auto'>"
        "<p>For full detail, open the original JPG in the Explorer and zoom in.</p>"
    )))


def preview_spreadsheet(path=SPREADSHEET):
    # Never bake an instructor/student spreadsheet into the published notebooks.
    if os.environ.get("WORKSHOP_DEMO_ANSWERS") or not path.is_file():
        return display(HTML(readable_html(
            "<b>Your Excel preview will appear here.</b> First create and run "
            "<code>_for_STUDENT/tasks/06_receipt_ocr/receipt_reader.py</code>, then rerun this cell. "
            "Expected output: <code>_for_STUDENT/outputs/06_receipt_ocr/first_receipt.xlsx</code>."
        )))
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            rows = list(islice(workbook.active.iter_rows(values_only=True), 11))
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, BadZipFile) as exc:
        return display(HTML(readable_html(f"<b>Could not open the spreadsheet.</b> {escape(str(exc))}")))
    table = "<table>" + "".join(
        "<tr>" + "".join(
            f"<{'th' if index == 0 else 'td'}>{escape('' if value is None else str(value))}</{'th' if index == 0 else 'td'}>"
            for value in row
        ) + "</tr>" for index, row in enumerate(rows)
    ) + "</table>"
    return display(HTML(readable_html("<b>Your spreadsheet · first 10 data rows</b>" + table)))
