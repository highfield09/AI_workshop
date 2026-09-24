from pathlib import Path

import pytest

from llm_workshop import receipt_support
from llm_workshop.course import ROOT
from scripts.build_course_notebook import build_workbooks
from scripts.check_receipt_assets import check_assets
from scripts.prepare_receipt_data import EXPECTED_IMAGES, REMOTE_IMAGES, REMOTE_CSV, prepare


def test_ocr_lesson_has_four_iterations_and_traceable_outputs():
    book = build_workbooks()["06_receipt_ocr"]
    source = "\n".join(c.source for c in book.cells)
    assert "osama hosam Abdellatif" in source
    assert "version 3" in source and "Database: Open Database" in source
    assert "first ten images" in source
    assert "seven item rows" in source and "one receipt row" in source
    for text in ['Iteration 1', 'Iteration 2', 'Iteration 3', 'Iteration 4', 'first 5', 'quantity', 'receipts_5.xlsx', 'batch_run.csv', 'item_quantities.png', 'histogram', 'not PDFs', 'Text Recognition:', 'max_new_tokens=4096']:
        assert text in source
    assert "receipt_reader.py" in source and "first_receipt.xlsx" in source
    assert "<details>" in source and "Hint · Build your prompt" in source
    assert "the whole folder yet" in source
    assert "Do not let the script copy answers from" in source
    assert "from llm_workshop.receipt_support" in book.cells[1].source
    assert source.index("Dataset, licensing and credit") < source.index("show_receipt_preview()")


def test_model_and_dataset_manifest_are_valid():
    check_assets()


def test_downloader_requests_only_explicit_first_subbatch_files(tmp_path, monkeypatch):
    import kagglehub
    cache = tmp_path / "cache"
    cache.mkdir()
    calls = []
    def fake_download(handle, *, path):
        assert "/versions/3" in handle
        calls.append(path)
        file = cache / Path(path).name
        file.write_bytes(b"test fixture")
        return str(file)
    monkeypatch.setattr(kagglehub, "dataset_download", fake_download)
    dest = tmp_path / "data"
    prepare(dest)
    expected = {REMOTE_CSV} | {f"{REMOTE_IMAGES}/{name}" for name in EXPECTED_IMAGES}
    assert set(calls) == expected and len(calls) == 11
    calls.clear()
    prepare(dest)
    assert calls == []  # Reuse all verified local files.
    older = dest / "batch1_1/batch1-0011.jpg"
    older.write_bytes(b"older download")
    prepare(dest)
    assert older.read_bytes() == b"older download" and calls == []
    (dest / "batch1_1/batch1-0001.jpg").write_bytes(b"changed")
    with pytest.raises(RuntimeError, match="differs"):
        prepare(dest)
    assert (dest / "batch1_1/batch1-0001.jpg").read_bytes() == b"changed"


def test_excel_preview_is_safe_and_demo_does_not_read_personal_data(tmp_path, monkeypatch):
    from openpyxl import Workbook
    path = tmp_path / "first_receipt.xlsx"
    workbook = Workbook()
    workbook.active.append(["seller_name", "grand_total"])
    workbook.active.append(["<script>PRIVATE TEST VALUE</script>", 6204.19])
    workbook.save(path)
    shown = []
    monkeypatch.setattr(receipt_support, "display", lambda content: shown.append(content.data))
    monkeypatch.setenv("WORKSHOP_DEMO_ANSWERS", str(tmp_path))
    receipt_support.preview_spreadsheet(path)
    assert "PRIVATE TEST VALUE" not in shown[-1]
    monkeypatch.delenv("WORKSHOP_DEMO_ANSWERS")
    receipt_support.preview_spreadsheet(path)
    assert "6204.19" in shown[-1]
    assert "&lt;script&gt;PRIVATE TEST VALUE&lt;/script&gt;" in shown[-1]
    assert "<script>PRIVATE TEST VALUE" not in shown[-1]
    path.write_bytes(b"not an Excel file")
    receipt_support.preview_spreadsheet(path)
    assert "Could not open the spreadsheet" in shown[-1]


def test_preview_handles_missing_receipt_without_downloading(tmp_path, monkeypatch):
    shown = []
    monkeypatch.setattr(receipt_support, "display", lambda content: shown.append(content.data))
    receipt_support.show_receipt_preview(tmp_path / "missing.jpg")
    assert "prepare_receipt_data.py" in shown[-1]


def test_glm_lesson_documents_local_runtime_and_validation():
    source = '\n'.join(c.source for c in build_workbooks()['06_receipt_ocr'].cells)
    for required in ['GLM-OCR', '16 GB RAM', 'Workbook 1 · Master setup', 'model.glm_reader', 'JSON schema', 'Tax Id']:
        assert required in source
    assert 'tesserocr' not in source and 'Tesseract' not in source


def test_excel_preview_can_select_each_sheet(tmp_path, monkeypatch):
    from openpyxl import Workbook
    path = tmp_path / 'first_receipt.xlsx'
    book = Workbook()
    book.active.title = 'Items'
    book.active.append(['quantity'])
    book.active.append([3])
    book.create_sheet('Receipts').append(['grand_total'])
    book.save(path)
    shown = []
    monkeypatch.delenv('WORKSHOP_DEMO_ANSWERS', raising=False)
    monkeypatch.setattr(receipt_support, 'display', lambda content: shown.append(content.data))
    receipt_support.preview_spreadsheet(path, sheet_name='Items')
    assert 'quantity' in shown[-1] and 'Items' in shown[-1]
    receipt_support.preview_spreadsheet(path, sheet_name='Receipts')
    assert 'grand_total' in shown[-1] and 'Receipts' in shown[-1]
    receipt_support.preview_spreadsheet(path, sheet_name='Missing')
    assert 'Could not open' in shown[-1]


def test_short_lesson_has_copyable_requirements_and_correct_viewer():
    book = build_workbooks()["06_receipt_ocr"]
    prompts = [c.source for c in book.cells if "copyable_prompt(" in c.source]
    assert len(prompts) == 2
    assert "Input:" in prompts[0] and "raw/batch1-0001.txt" in prompts[0]
    assert "unit_net_price" in prompts[1] and "Receipts" in prompts[1]
    source = "\n".join(c.source for c in book.cells)
    assert "GrapeCity.gc-excelviewer" in source
    assert "Illustrative output—not a live download" in source
    assert "receipts_20" not in source and "first 20" not in source
    assert "model.glm_reader" in source


def test_ocr_preparation_is_at_the_beginning_of_workbook_one():
    books = build_workbooks()
    opening = books["01_start_here"].cells[0].source
    assert "Master setup" in opening
    first = "\n".join(c.source for c in books["01_start_here"].cells)
    assert first.index("Master setup") < first.index("START HERE")
    for required in ["setup_glm_ocr.sh", "prepare_receipt_data.py",
                     "check_receipt_assets.py --require-data", "16 GB RAM",
                     "Run All", "does **not** load"]:
        assert required in first
    sixth = "\n".join(c.source for c in books["06_receipt_ocr"].cells)
    assert "sh _for_TRAINER/scripts/setup_glm_ocr.sh" not in sixth
    assert "(01_start_here.ipynb)" in sixth


def test_receipt_preview_follows_the_task_immediately():
    cells = build_workbooks()["06_receipt_ocr"].cells
    task = next(i for i,c in enumerate(cells) if "TASK · READ THE RECEIPT" in c.source)
    assert "preview directly below" in cells[task].source
    assert cells[task+1].source == "show_receipt_preview()"
    assert "copyable_prompt(" in cells[task+2].source
    assert sum(c.source == "show_receipt_preview()" for c in cells) == 1


def test_startup_example_has_theme_independent_contrast_and_runtime_guidance():
    cells = build_workbooks()["06_receipt_ocr"].cells
    example = next(c.source for c in cells if "What model startup can look like" in c.source)
    assert '<pre style="background:#F8FAFC !important;color:#1D2939 !important;' in example
    assert 'color:#1D2939 !important">Loading local GLM-OCR processor' in example
    source = "\n".join(c.source for c in cells)
    assert ".venv/bin/python _for_STUDENT/tasks/06_receipt_ocr/receipt_reader.py" in source
    assert "TASK_DIR.parents[1]" in source
    assert "Choosing a notebook kernel does not change" in source
